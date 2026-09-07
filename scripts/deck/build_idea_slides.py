"""Fill the OFFICIAL SIH 2026 idea-submission template for SIH26066.

Base: scripts/deck/template/SIH2026-IDEA-Presentation-Format.pptx -- the file downloaded from the
SIH portal, copied in unmodified so the build is reproducible. Its own chrome (the SIH logo, the
"SMART INDIA HACKATHON 2026" title page, the footer bar, the slide numbers, the team-name oval) is
left exactly as the portal ships it. We fill it; we do not redraw it.

Output: OceanEmbed_SIH26066_IDEA.pptx -- six slides, which is the template's stated maximum.

WHAT THE TEMPLATE'S OWN INSTRUCTION SLIDE REQUIRES (slide 7 of the original)
  * maximum six slides including the title page  -> slide 7 is deleted here
  * points, diagrams and infographics, not paragraphs
  * upload as PDF, not PPTX  -> export from PowerPoint; there is no Office renderer on this machine

EVERY MEASURED NUMBER IS READ FROM artifacts/frozen_manifest.json AND THE METRICS ARTIFACT IT
NAMES, never typed here, so a re-freeze moves the slide too. _guard() refuses to save a deck that
quotes a retracted result or the GLORYS-fed comparator.

Design spec: docs/superpowers/specs/2026-09-07-sih-idea-deck-design.md
Run:  python scripts/deck/build_idea_slides.py
"""
from __future__ import annotations

import copy
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

import diagrams

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template",
                        "SIH2026-IDEA-Presentation-Format.pptx")
MANIFEST = os.path.join(REPO, "artifacts", "frozen_manifest.json")
DECK = os.path.join(REPO, "OceanEmbed_SIH26066_IDEA.pptx")

# ---------------------------------------------------------------- portal details
#: [VERIFIED] SIH26066; theme/category/organisation from the PS transcription confirmed in
#: docs/ARJHUN_EXECUTION_PLAN.md:30 ("Theme: Disaster Management. Category: Software.
#: Organization: INCOIS, MoES"). README.md's "Space Technology" is an older informal note.
#: [UNKNOWN] the portal's verbatim PS title, the team ID and the team name are recorded nowhere in
#: this repo. They carry a visible marker so they cannot reach a jury unfilled.
MARK = "‹fill from portal›"
PS_ID = "SIH26066"
PS_TITLE = f"Subsurface ocean temperature from satellite observations  {MARK}"
THEME = "Disaster Management"
PS_CATEGORY = "Software"
ORG = "Ministry of Earth Sciences (INCOIS)"
TEAM_ID = MARK
TEAM_NAME = MARK
PROJECT = "OceanEmbed"

#: The template's content band: the title placeholder ends at y=1.20, the footer bar starts at
#: y=6.95, and the SIH logo occupies x>10.70 above y=1.16.
TOP, BOTTOM = 1.26, 6.88

# ---------------------------------------------------------------- palette
INK = RGBColor.from_string("12232B")
HEAD_BLUE = RGBColor.from_string("1F4E79")
GREY = RGBColor.from_string("6B7C87")
TEAL = RGBColor.from_string("1D7874")

BODY_FONT = "Arial"          # the template's own body font
SYM_FONT = "Segoe UI Symbol"  # has ❖ (U+2756) and ➤ (U+27A4)

FORBIDDEN = {
    "0.9267": "pre-embargo T_SEQ 31 leg (INVALID_PRE_EMBARGO.md, 'never quote')",
    "0.8529": "pre-embargo T_SEQ leg, checkpoint overwritten",
    "0.9096": "pre-embargo T_SEQ leg, checkpoint overwritten",
    "0.8873": "tscast_stage1_embargo_withUV_s42 -- the GLORYS comparator, not the deliverable",
    "0.8548": "stage-2 GLORYS run -- reanalysis inputs, fails the PS satellite-only clause",
}


# ================================================================ numbers
def load_numbers() -> dict[str, str]:
    """Read every quotable figure straight out of the freeze."""
    with open(MANIFEST, encoding="utf-8") as f:
        m = json.load(f)
    c = m["claims"][m["deliverable_key"]]
    assert c["deliverable"] is True, "manifest deliverable_key does not point at the deliverable"
    assert c["input_source"] == "satellite", "the deliverable must be the satellite-input model"

    with open(os.path.join(REPO, "artifacts", c["metrics_file"]), encoding="utf-8") as f:
        per_all = json.load(f)
    per = per_all["metrics"]
    depths, rmse, clim = per["depths_m"], per["rmse"], per["rmse_climatology"]
    wins = sum(1 for r, k in zip(rmse, clim) if r < k)
    band = [(d, r) for d, r in zip(depths, rmse) if 50 <= d <= 150]
    peak_d, peak_r = max(band, key=lambda t: t[1])

    return {
        "rmse": f"{c['overall_rmse']:.2f}",
        "bias": f"+{c['overall_bias']:.2f}",
        "corr": f"{c['overall_correlation']:.2f}",
        "skill": f"{c['overall_skill_rmse_ratio'] * 100:.0f}%",
        "n": f"{c['overall_n']:,}",
        "profiles": f"{c['argo_profiles']:,}",
        "depths": str(c["n_depths"]),
        "tseq": str(c["T_SEQ"]),
        "channels": str(len(c["channels"])),
        "leak": f"{c['selection_leak']['measured_cost_degC']:.2f}",
        "wins": str(wins),
        "peak_rmse": f"{peak_r:.2f}",
        "peak_depth": str(peak_d),
        "band_lo": f"{min(r for _, r in band):.2f}",
        "params": f"{round((per_all['n_params_encoder'] + per_all['n_params_decoder']) / 1000)}k",
        "train_min": f"{per_all['train_seconds'] / 60:.0f}",
        "device": "GPU" if per_all["device"] == "cuda" else per_all["device"].upper(),
    }


def _collect_test_count() -> str | None:
    """How many tests the suite collects, right now. A hardcoded count is exactly what this repo
    keeps getting burned by -- the inherited "530 passing" was already 500 short."""
    import subprocess

    py = os.path.join(REPO, ".venv", "Scripts", "python.exe")
    if not os.path.exists(py):
        py = sys.executable
    try:
        out = subprocess.run([py, "-m", "pytest", "tests", "-q", "--collect-only"],
                             cwd=REPO, capture_output=True, text=True, timeout=180).stdout
    except Exception:
        return None
    hit = re.search(r"(\d+) tests? collected", out)
    return f"{int(hit.group(1)):,}" if hit else None


N = load_numbers()
_tc = _collect_test_count()
TESTS = f"{_tc} tests" if _tc else None

# ================================================================ content
SOLUTION = [
    ("Comprehensive OceanEmbed Subsurface Platform:",
     f"A web dashboard and reproducible pipeline that reconstructs daily ocean temperature at "
     f"{N['depths']} depths from 0 to 1000 m across the North Indian Ocean, from surface "
     f"observations alone."),
    ("Satellite-Only Input:",
     "Runs on OSTIA sea-surface temperature, DUACS altimetry, SMOS-blended salinity, GLOBCURRENT "
     "surface currents and observational winds — no in-situ float is needed at prediction time, "
     "which is what makes it operational every day."),
    ("TS-Cast-NIO Model:",
     f"A 3-D CNN compresses an {N['tseq']}-day window of {N['channels']} surface maps into a "
     f"128-number ocean embedding, and a decoder expands it into a full profile — {N['params']} "
     f"parameters, kept deliberately small for the data available."),
    ("Uncertainty on Every Value:",
     "The model predicts a spread alongside each temperature, so the product states where it "
     "should not be trusted instead of showing one confident number everywhere."),
    ("Independent Validation:",
     f"Scored against {N['profiles']} Argo float profiles never seen in training — {N['rmse']} °C "
     f"RMSE over {N['n']} comparisons, {N['skill']} better than climatology, and better at "
     f"{N['wins']} of the {N['depths']} depths."),
    ("Derived Ocean Products:",
     "Mixed-layer depth, thermocline depth, cyclone heat potential (TCHP / D26), eddy and "
     "upwelling detection, and anomaly maps — computed from the reconstructed profile, never "
     "assumed."),
    ("Observation-Priority Map:",
     "The uncertainty field ranks where the ocean is least known, turning a gap in knowledge into "
     "an answer to where the next Argo float should be deployed."),
]

STACK = [
    ("Languages & Core:", "Python 3.12, NumPy, pandas, PyArrow, PyYAML."),
    ("Ocean Data:", "xarray, netCDF4, SciPy, copernicusmarine (CMEMS), argopy, erddapy."),
    ("AI / ML:", f"PyTorch — 3-D CNN encoder + depth decoder, {N['params']} parameters; "
                 f"scikit-learn; LightGBM baseline and quantile uncertainty; CUDA."),
    ("Frontend & Visualisation:", "Streamlit — 15 feature modules behind one lazy registry; "
                                  "Plotly; Matplotlib."),
    ("API & Export:", "FastAPI, Uvicorn, and NetCDF / CSV / Parquet export."),
    ("Reproducibility:", f"pytest — {TESTS or 'a suite'} covering the code and the claims; Git; "
                         f"ruff; seeded configs; a SHA-256-verified frozen checkpoint manifest."),
]

FEAS_LEFT = [
    ("head", "Feasibility", "2E75B6"),
    ("1.", "Technical:", "Already built end-to-end — 388 gap-free days of satellite input, a "
                         "trained model, and a running instrument. Not a proposal."),
    ("2.", "Economic:", f"Every input is free and open. CMEMS and Argo cost nothing, training "
                        f"takes about {N['train_min']} minutes on one {N['device']}, and "
                        f"inference runs on a laptop."),
    ("3.", "Operational:", "The satellite L4 products are published daily and operationally, so "
                           "the pipeline runs forward in near-real-time with no in-situ input."),
    ("head", "Viability", "2E9E5B"),
    ("4.", "Scales beyond the pilot:", "The architecture is region-agnostic — retargeting to "
                                       "another basin is a data-bundle change, not a redesign."),
    ("5.", "Reproducible and auditable:", f"Seed, config and a SHA-256-verified checkpoint are "
                                          f"frozen in a manifest, with {TESTS or 'a test suite'} "
                                          f"that guard the claims as well as the code."),
    ("head", "Challenges", "C55A11"),
    ("6.", "The thermocline is the hard layer:", f"Error runs {N['band_lo']}–{N['peak_rmse']} °C "
                                                 f"through 50–150 m, peaking at {N['peak_depth']} "
                                                 f"m where temperature falls fastest."),
    ("7.", "Uncertainty is improved, not calibrated:", "Coverage runs below nominal in the mixed "
                                                       "layer. We report the shortfall rather "
                                                       "than widening the band to hide it."),
    ("8.", "Epoch selection was not independent:", f"The shipped epoch was chosen on the test "
                                                   f"period; a leak-free protocol costs "
                                                   f"+{N['leak']} °C over three seeds."),
    ("9.", "INCOIS gridded Argo is unreachable:", "Their data layer is down, so validation uses "
                                                  "individual argopy floats instead."),
    ("head", "Use Cases", "7030A0"),
    ("10.", "Cyclone intensity:", "upper-100 m heat content, the fuel a surface map cannot see."),
    ("11.", "Monsoon & fisheries:", "Bay of Bengal stratification and Arabian Sea upwelling."),
    ("12.", "Next-float targeting:", "an uncertainty-ranked map of where the ocean is least "
                                     "known."),
]

FEAS_RIGHT = [
    ("head", "Deployment Potential", "2E75B6"),
    ("1.", "Hand-off to INCOIS / IMD:", "A daily NetCDF field on a standard grid — directly "
                                        "ingestible by existing forecast and assimilation "
                                        "chains."),
    ("2.", "Zero licence cost:", "Open Copernicus and Argo data on an open-source stack."),
    ("3.", "Fisheries & naval advisories:", "The same profile drives potential-fishing-zone and "
                                            "sonar-range products agencies already issue."),
    ("head", "Strategies", "2E9E5B"),
    ("4.", "Robust data & protocol:", f"388 gap-free days, a 5-day embargo between train and "
                                      f"test, scored on {N['profiles']} unseen floats."),
    ("5.", "Modular integration:", "15 lazily-loaded feature modules behind one registry — a "
                                   "broken feature contains itself."),
    ("6.", "Refusal over invention:", "Below the seafloor the model returns nothing rather than a "
                                      "number, and the scorer declines those depths too."),
]

AUDIENCE = [
    ("INCOIS & IMD:", "A daily subsurface field with an uncertainty band — an input their cyclone "
                      "and ocean-state forecasts currently lack at this cadence."),
    ("Fisheries & coastal communities:", "Thermocline and mixed-layer depth drive fishing zones; "
                                         "a daily profile sharpens advisories that today rest on "
                                         "surface temperature alone."),
    ("Indian Navy & shipping:", "Sound-speed structure follows the temperature profile, feeding "
                                "sonar-range and routing decisions."),
    ("Ocean researchers & Argo:", "A quantitative answer to where the next float is worth "
                                  "deploying."),
    ("Disaster management:", "Upper-ocean heat is the fuel term in rapid cyclone intensification "
                             "— the failure mode that costs the most lives on this coast."),
]

BENEFITS = [
    ("Social:", "6E8B5E", [
        "Better cyclone-intensity input for the most densely populated cyclone basin.",
        "Sharper fishing-zone advisories for small-scale fishers.",
        "Uncertainty shown on screen, so an advisory states its confidence instead of implying "
        "certainty.",
    ]),
    ("Technological:", "C0504D", [
        "Targets the North Indian Ocean, where most published work targets the Pacific.",
        "An uncertainty head trained jointly with temperature, not bolted on afterwards.",
        "Negative results published beside positive ones — the FiLM decoder and the density loss "
        "were both measured harmful and dropped.",
    ]),
    ("Economic:", "7C61A8", [
        "Zero data-licensing cost: Copernicus and Argo are open, and so is the stack.",
        f"About {N['train_min']} minutes of {N['device']} training and laptop-class inference.",
        "Targets Argo deployment — the costliest part of ocean observation — where information is "
        "scarcest.",
    ]),
    ("Environmental:", "3E8E4F", [
        "Monitors ocean heat content, the dominant store of excess planetary heat.",
        "Sees subsurface heat and upwelling that surface-only monitoring is blind to.",
        "Reduces redundant ship survey by targeting where observation adds information.",
    ]),
]

PAPERS = [
    ("Meng et al., 2021.", "Reconstruction of Three-Dimensional Temperature and Salinity Fields "
                           "From Satellite Observations. JGR: Oceans. doi:10.1029/2021JC017605"),
    ("TS-Cast, 2026.", "Deep Learning for Subsurface Ocean Reconstruction from Satellite "
                       "Observations in the Northwestern Pacific. Ocean Science 22, 2161."),
    ("DORS, 2022.", "Subsurface Temperature Reconstruction for the Global Ocean 1993-2020 Using "
                    "Satellite Observations and Deep Learning. Remote Sensing 14(13), 3198. "
                    "doi:10.3390/rs14133198"),
    ("FFPG-net, 2025.", "Feature fusion with physical guidance for subsurface thermal structure "
                        "reconstruction."),
    ("NeSPReSO, 2025.", "PCA + neural network subsurface prediction, Gulf of Mexico."),
]

SOURCES = [
    ("Surface inputs (Copernicus Marine):", "OSTIA sea-surface temperature; DUACS L4 altimetry "
                                            "(cmems_obs-sl_glo_phy-ssh); SMOS-blended salinity "
                                            "(cmems_obs-mob_glo_phy-sss); GLOBCURRENT currents "
                                            "(cmems_obs-mob_glo_phy-cur); L4 wind "
                                            "(cmems_obs-wind_glo_phy)."),
    ("Training target:", "GLORYS12V1 reanalysis, cmems_mod_glo_phy_my_0.083deg_P1D "
                         "(doi:10.48670/moi-00021), as the problem statement specifies."),
    ("Independent validation:", f"Argo float profiles via argopy — {N['profiles']} profiles, "
                                f"{N['n']} depth comparisons, never seen in training."),
    ("Problem statement:", f"{PS_ID}, {ORG} — three-dimensional ocean temperature from only "
                           f"surface satellite observations; 5–30°N, 45–105°E, 0.25°, daily."),
]

CAVEAT = ("Citation discipline: the five papers above are reviewed at abstract level "
          "(docs/LITERATURE_MATRIX.md). They are cited for existence and general approach only. "
          "No numerical comparison against their published results appears anywhere in this deck.")


# ================================================================ pptx helpers
def _drop(shape):
    """Remove a shape from its slide."""
    shape._element.getparent().remove(shape._element)


def _find(slide, *prefixes):
    return [sh for sh in slide.shapes if sh.name.startswith(prefixes)]


def _retext(shape, text):
    """Replace a shape's text, keeping the template's own run formatting."""
    tf = shape.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run().text = text


def _tb(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _p(tf, first, *, space_after=6, space_before=0, indent=0.0, line=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    if indent:
        emu = int(Inches(indent))
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(emu))
        pPr.set("indent", str(-emu))
    return p


def _r(p, text, *, size, bold=False, color=INK, font=BODY_FONT, underline=False):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.underline = underline
    return run


def _section(slide, x, y, text, *, size=22, w=8.0):
    tf = _tb(slide, x, y, w, 0.42)
    p = _p(tf, True, space_after=0)
    _r(p, "❖", size=size, bold=True, color=HEAD_BLUE, font=SYM_FONT)
    _r(p, " " + text, size=size, bold=True, color=HEAD_BLUE, underline=True)


def _bullets(tf, items, *, size, marker="➤", indent=0.28, space=7, lead_size=None):
    for i, (lead, rest) in enumerate(items):
        p = _p(tf, i == 0, space_after=space, indent=indent)
        _r(p, marker + "  ", size=size * 0.85, color=INK, font=SYM_FONT)
        _r(p, lead + " ", size=lead_size or size, bold=True)
        _r(p, rest, size=size)


def _numbered(tf, rows, *, size, space=4):
    first = True
    for kind, a, b in rows:
        if kind == "head":
            p = _p(tf, first, space_after=3, space_before=0 if first else 7)
            _r(p, "■ ", size=size + 1.5, color=RGBColor.from_string(b), font=SYM_FONT)
            _r(p, a, size=size + 1.5, bold=True, color=RGBColor.from_string(b))
        else:
            p = _p(tf, first, space_after=space, indent=0.26)
            _r(p, kind + " ", size=size, bold=True, color=GREY)
            _r(p, a + " ", size=size, bold=True)
            _r(p, b, size=size)
        first = False


def _pic(slide, name, x, y, w):
    return slide.shapes.add_picture(os.path.join(REPO, "artifacts", "deck", name),
                                    Inches(x), Inches(y), width=Inches(w))


# ================================================================ slides
def fill_title(slide):
    """The portal's own fields, on the template's title page."""
    box = next(sh for sh in slide.shapes if sh.name == "TextBox 9")
    tf = box.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in tf.paragraphs[0].runs:
        r._r.getparent().remove(r._r)

    rows = [("Problem Statement ID – ", PS_ID),
            ("Problem Statement Title – ", PS_TITLE),
            ("Theme – ", THEME),
            ("PS Category – ", PS_CATEGORY),
            ("Organisation – ", ORG),
            ("Team ID – ", TEAM_ID),
            ("Team Name – ", TEAM_NAME)]
    for i, (label, value) in enumerate(rows):
        p = _p(tf, i == 0, space_after=7)
        _r(p, label, size=15, bold=True)
        _r(p, value, size=15)


def fill_solution(slide):
    _retext(next(sh for sh in slide.shapes if sh.name == "Title 1"), PROJECT)
    _section(slide, 0.45, TOP, "Proposed Solution", size=22, w=7.2)
    _bullets(_tb(slide, 0.55, TOP + 0.50, 7.10, BOTTOM - TOP - 0.50), SOLUTION, size=10.5)
    _pic(slide, "arch.png", 7.95, TOP + 0.06, 4.55)


def fill_technical(slide):
    _section(slide, 0.45, TOP, "Technology Stack", size=22, w=7.0)
    tf = _tb(slide, 0.55, TOP + 0.50, 6.65, 3.90)
    for i, (lead, rest) in enumerate(STACK):
        p = _p(tf, i == 0, space_after=9)
        _r(p, lead + " ", size=11, bold=True)
        _r(p, rest, size=11)
    _pic(slide, "chips.png", 7.55, TOP + 0.02, 5.15)

    cap = _tb(slide, 0.46, 5.90, 12.40, 0.22)
    p = _p(cap, True, space_after=0)
    _r(p, "METHODOLOGY AND PROCESS FOR IMPLEMENTATION", size=9, bold=True, color=GREY)
    _pic(slide, "flow.png", 0.46, 6.13, 12.40)


def fill_feasibility(slide):
    _numbered(_tb(slide, 0.40, TOP, 6.15, BOTTOM - TOP), FEAS_LEFT, size=8.6)
    _numbered(_tb(slide, 6.75, TOP, 6.15, 3.25), FEAS_RIGHT, size=8.6)
    _pic(slide, "facts.png", 6.75, 4.63, 6.15)


def fill_impact(slide):
    _pic(slide, "tree.png", 0.42, TOP + 0.04, 6.05)
    _section(slide, 0.42, 3.78, "Potential impact on the target audience", size=14, w=6.3)
    _bullets(_tb(slide, 0.46, 4.22, 6.05, BOTTOM - 4.22), AUDIENCE, size=9,
             marker="•", indent=0.18, space=5)

    _section(slide, 6.75, TOP, "Benefits of the solution", size=18, w=6.2)
    tf = _tb(slide, 6.80, TOP + 0.46, 6.15, BOTTOM - TOP - 0.46)
    first = True
    for head, col, lines in BENEFITS:
        p = _p(tf, first, space_after=3, space_before=0 if first else 6)
        _r(p, head, size=11, bold=True, color=RGBColor.from_string(col))
        first = False
        for line in lines:
            q = _p(tf, False, space_after=2, indent=0.16)
            _r(q, "•  ", size=9, color=GREY)
            _r(q, line, size=9)


def fill_references(slide):
    _section(slide, 0.45, TOP, "Key literature", size=17, w=6.0)
    tf = _tb(slide, 0.55, TOP + 0.44, 6.20, 3.80)
    for i, (lead, rest) in enumerate(PAPERS):
        p = _p(tf, i == 0, space_after=8, indent=0.20)
        _r(p, "▸  ", size=9.5, color=TEAL, font=SYM_FONT)
        _r(p, lead + " ", size=9.5, bold=True)
        _r(p, rest, size=9.5)

    _section(slide, 7.00, TOP, "Data and problem statement", size=17, w=6.0)
    tf = _tb(slide, 7.10, TOP + 0.44, 5.85, 3.80)
    for i, (lead, rest) in enumerate(SOURCES):
        p = _p(tf, i == 0, space_after=8, indent=0.20)
        _r(p, "▸  ", size=9.5, color=TEAL, font=SYM_FONT)
        _r(p, lead + " ", size=9.5, bold=True)
        _r(p, rest, size=9.5)

    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(5.95),
                                  Inches(12.40), Inches(0.85))
    note.fill.solid()
    note.fill.fore_color.rgb = RGBColor.from_string("FFF6E8")
    note.line.color.rgb = RGBColor.from_string("E0913A")
    note.shadow.inherit = False
    tf = note.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.14)
    p = _p(tf, True, space_after=0)
    _r(p, CAVEAT, size=9.5, color=RGBColor.from_string("6B4415"))


# ================================================================ guard + build
def _guard(prs):
    blob = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                     if sh.has_text_frame)

    for bad, why in FORBIDDEN.items():
        if bad in blob:
            raise AssertionError(f"deck quotes {bad} -- {why}")
    for key, val in (("rmse", N["rmse"]), ("profiles", N["profiles"]), ("n", N["n"])):
        if val not in blob:
            raise AssertionError(f"deck never states the manifest {key} ({val}) -- guard vacuous")
    if f"{N['wins']} of the {N['depths']} depths" not in blob:
        raise AssertionError("the depth-win qualifier is missing or stale")
    if int(N["wins"]) >= int(N["depths"]):
        raise AssertionError("climatology beaten at every depth -- verify before claiming it")
    if len(prs.slides) != 6:
        raise AssertionError(f"{len(prs.slides)} slides; the template's own instruction is a "
                             f"maximum of six including the title page")
    for pointer in ("Detailed explanation of the proposed solution",
                    "Technologies to be used", "Potential challenges and risks"):
        if pointer in blob:
            raise AssertionError(f"the template's prompt text {pointer!r} is still on a slide")
    return blob


def main() -> int:
    diagrams.build_all()

    prs = Presentation(TEMPLATE)
    assert len(prs.slides) == 7, f"template has {len(prs.slides)} slides, expected 7"

    # The template ships each content slide with its prompt text in "TextBox 8" and a
    # "Your Team Name" oval. Drop the prompts, fill the ovals.
    for slide in list(prs.slides)[1:6]:
        for sh in _find(slide, "TextBox"):
            _drop(sh)
        for sh in _find(slide, "Oval"):
            _retext(sh, TEAM_NAME)

    slides = list(prs.slides)
    fill_title(slides[0])
    fill_solution(slides[1])
    fill_technical(slides[2])
    fill_feasibility(slides[3])
    fill_impact(slides[4])
    fill_references(slides[5])

    # Slide 7 is the template's own "IMPORTANT INSTRUCTIONS" page, which it says to delete before
    # uploading, and which would breach its own six-slide maximum.
    xml_slides = prs.slides._sldIdLst
    last = list(xml_slides)[-1]
    prs.part.drop_rel(last.rId)
    xml_slides.remove(last)

    blob = _guard(prs)
    prs.save(DECK)

    print(f"wrote {DECK}")
    print(f"  {len(prs.slides)} slides  ·  {os.path.getsize(DECK):,} bytes  "
          f"·  template chrome preserved, instruction slide removed")
    print(f"  from the freeze: RMSE {N['rmse']} °C, skill {N['skill']}, {N['profiles']} profiles, "
          f"n={N['n']}, {N['wins']}/{N['depths']} depths")
    print(f"  {len(blob):,} characters of slide text, no forbidden figure present")
    print(f"  STILL TO FILL ({MARK}): PS title, team ID, team name")
    print("  UPLOAD AS PDF -- the template requires it; export from PowerPoint")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
